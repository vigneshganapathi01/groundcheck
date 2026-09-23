"""GroundCheck — local web dashboard for LettuceDetect hallucination detection.

Run:
    .venv/bin/python app/server.py            # then open http://127.0.0.1:8765
"""

from __future__ import annotations

import os
import threading
import time
from pathlib import Path

os.environ.setdefault("TRANSFORMERS_VERBOSITY", "error")
os.environ.setdefault("HF_HUB_DISABLE_PROGRESS_BARS", "1")

from fastapi import FastAPI, File, HTTPException, UploadFile
from fastapi.responses import FileResponse
from pydantic import BaseModel, Field

from lettucedetect.detectors.prompt_utils import PromptUtils
from lettucedetect.models.inference import HallucinationDetector

STATIC_DIR = Path(__file__).parent / "static"

MODELS = {
    "base": {
        "id": "KRLabsOrg/lettucedect-base-modernbert-en-v1",
        "label": "ModernBERT base · 150M · English",
        "size": "~600 MB download",
    },
    "large": {
        "id": "KRLabsOrg/lettucedect-large-modernbert-en-v1",
        "label": "ModernBERT large · 395M · English",
        "size": "~1.6 GB download",
    },
}

_detectors: dict[str, HallucinationDetector] = {}
_lock = threading.Lock()


def get_detector(key: str) -> HallucinationDetector:
    """Load a detector once and keep it in memory for later requests."""
    with _lock:
        if key not in _detectors:
            _detectors[key] = HallucinationDetector(
                method="transformer", model_path=MODELS[key]["id"]
            )
        return _detectors[key]


class DetectRequest(BaseModel):
    context: list[str] = Field(..., min_length=1)
    question: str | None = None
    answer: str = Field(..., min_length=1)
    model: str = "base"
    threshold: float = Field(0.5, ge=0.05, le=0.99)
    # "sentence": score each sentence of the answer on its own (default; see score_answer)
    # "whole": score the full answer in one pass, as LettuceDetect does
    mode: str = "sentence"


app = FastAPI(title="GroundCheck")

# Changes whenever the server restarts; open pages compare it to know they're outdated.
BUILD_ID = str(int(time.time()))


@app.get("/")
def index() -> FileResponse:
    # no-store: a reload always gets the current page, never a cached older version.
    return FileResponse(STATIC_DIR / "index.html", headers={"Cache-Control": "no-store"})


@app.get("/api/version")
def version() -> dict:
    return {"build": BUILD_ID}


def unreadable_reason(text: str) -> str | None:
    """Return why text looks like binary garbage (e.g. a .docx read as plain text), else None."""
    sample = text[:20000]
    if not sample.strip():
        return None
    if sample.startswith("PK\x03\x04") or "word/document.xml" in sample or "[Content_Types].xml" in sample:
        return "it is the raw bytes of a .docx file"
    if sample.startswith("%PDF-"):
        return "it is the raw bytes of a PDF file"
    bad = sum(ch == "\ufffd" or (ord(ch) < 32 and ch not in "\n\r\t") for ch in sample)
    if bad / len(sample) > 0.02:
        return f"{bad / len(sample):.0%} of it is unreadable characters"
    return None


@app.get("/api/models")
def list_models() -> dict:
    return {
        key: {**meta, "loaded": key in _detectors} for key, meta in MODELS.items()
    }


def _token_probs(inner, context: list[str], question: str | None, text: str) -> tuple[list[float], int]:
    """P(hallucinated) for each token of text (plus trailing [SEP]) against the context.

    Long contexts are split into chunks that fit the model. LettuceDetect combines chunks
    with max(), which flags any fact that appears in only one chunk; here a token counts
    as supported if ANY chunk supports it (min).
    """
    groups = inner._group_passages_into_chunks(context, question, text)
    runs = [
        inner._predict_single(PromptUtils.format_context(g, question, inner.lang), text, "tokens")
        for g in groups
    ]
    return [min(run[i]["prob"] for run in runs if i < len(run)) for i in range(len(runs[0]))], len(groups)


def _sentences(text: str) -> list[tuple[int, int]]:
    """Character ranges of the sentences (and separate lines) in text."""
    import re

    ranges, start = [], 0
    for m in re.finditer(r"(?<=[.!?])\s+|\n+", text):
        ranges.append((start, m.start()))
        start = m.end()
    ranges.append((start, len(text)))
    return [(a, b) for a, b in ranges if text[a:b].strip()]


def score_answer(inner, context, question, answer, mode) -> tuple[list[dict], int]:
    """Score every token of the answer; returns (tokens with char offsets, chunks per pass).

    In "sentence" mode each sentence is scored on its own. With the whole answer in one
    pass, correct sentences dilute the signal on a wrong one: in testing, "within 30
    calendar days" (the policy says 15) peaked at 0.33 inside the full answer but 0.68-0.97
    on its own, while correct sentences stayed near 0.00 either way.
    """
    pieces = _sentences(answer) if mode == "sentence" else [(0, len(answer))]
    tokens, n_chunks = [], 1
    for a, b in pieces:
        text = answer[a:b]
        probs, n_chunks = _token_probs(inner, context, question, text)
        offsets = inner.tokenizer(text, add_special_tokens=False, return_offsets_mapping=True)["offset_mapping"]
        for p, (s, e) in zip(probs, offsets):
            if s == e:
                continue
            s, e = a + s, a + e
            if tokens and s < tokens[-1]["end"]:
                # Byte-level BPE splits some characters (e.g. ₹) into several tokens that
                # share one character range; merge them so the character appears once.
                last = tokens[-1]
                last["end"] = max(last["end"], e)
                last["text"] = answer[last["start"] : last["end"]]
                last["prob"] = max(last["prob"], round(p, 4))
                continue
            tokens.append({"text": answer[s:e], "start": s, "end": e, "prob": round(p, 4)})
    return tokens, n_chunks


@app.post("/api/detect")
def detect(req: DetectRequest) -> dict:
    if req.model not in MODELS:
        raise HTTPException(400, f"Unknown model '{req.model}'")
    context = [p.strip() for p in req.context if p.strip()]
    if not context:
        raise HTTPException(400, "Provide at least one non-empty context passage.")
    if not req.answer.strip():
        raise HTTPException(400, "The answer to check is empty.")
    question = (req.question or "").strip() or None
    for i, passage in enumerate(context, 1):
        if reason := unreadable_reason(passage):
            raise HTTPException(
                422,
                f"Source passage {i} isn't readable text ({reason}). Checking against it would give "
                "meaningless results. Remove it and add the file with the Upload button, which extracts "
                "the text properly. If you already did, reload the page (Cmd+R) to get the latest version.",
            )
    if reason := unreadable_reason(req.answer):
        raise HTTPException(422, f"The AI response isn't readable text ({reason}). Use the upload icon to load it from a file.")

    t0 = time.perf_counter()
    was_loaded = req.model in _detectors
    detector = get_detector(req.model)
    load_s = 0.0 if was_loaded else time.perf_counter() - t0

    inner = detector.detector
    answer_tokens = len(inner.tokenizer(req.answer, add_special_tokens=False)["input_ids"])
    if answer_tokens > inner.max_length - 512:
        raise HTTPException(
            413,
            f"The AI response is too long to check in one go ({answer_tokens:,} tokens; "
            f"the limit is about {inner.max_length - 512:,}). Split it into smaller parts.",
        )

    if req.mode not in ("sentence", "whole"):
        raise HTTPException(400, "mode must be 'sentence' or 'whole'")
    t1 = time.perf_counter()
    tokens, n_chunks = score_answer(inner, context, question, req.answer, req.mode)
    infer_s = time.perf_counter() - t1
    for tok in tokens:
        tok["flagged"] = tok["prob"] >= req.threshold

    # Merge neighbouring flagged tokens into spans, trimming leading whitespace.
    spans: list[dict] = []
    prev_flagged = False
    for tok in tokens:
        if tok["flagged"] and prev_flagged:
            spans[-1]["end"] = tok["end"]
            spans[-1]["confidence"] = max(spans[-1]["confidence"], tok["prob"])
        elif tok["flagged"]:
            spans.append({"start": tok["start"], "end": tok["end"], "confidence": tok["prob"]})
        prev_flagged = tok["flagged"]
    for s in spans:
        text = req.answer[s["start"] : s["end"]]
        s["start"] += len(text) - len(text.lstrip())
        s["text"] = req.answer[s["start"] : s["end"]]

    flagged_chars = sum(s["end"] - s["start"] for s in spans)
    answer_chars = len(req.answer.strip()) or 1
    return {
        "model": {"key": req.model, **MODELS[req.model]},
        "threshold": req.threshold,
        "verdict": "hallucinated" if spans else "supported",
        "spans": spans,
        "tokens": tokens,
        "stats": {
            "span_count": len(spans),
            "flagged_pct": round(100 * flagged_chars / answer_chars, 1),
            "max_confidence": max((t["prob"] for t in tokens), default=0.0),
            "mean_prob": round(sum(t["prob"] for t in tokens) / max(len(tokens), 1), 4),
            "token_count": len(tokens),
            "flagged_tokens": sum(t["flagged"] for t in tokens),
            "chunks": n_chunks,
            "mode": req.mode,
            "inference_ms": round(infer_s * 1000),
            "load_ms": round(load_s * 1000),
        },
    }


# ---------------------------------------------------------------------------
# File upload: extract text from PDF / DOCX / TXT / MD
# ---------------------------------------------------------------------------

MAX_UPLOAD_BYTES = 20 * 1024 * 1024
PASSAGE_CHARS = 1200  # keeps passages small enough to chunk cleanly


def _clean(text: str) -> str:
    """Undo PDF line wrapping: join hyphenated breaks and single newlines, keep paragraphs."""
    import re
    import unicodedata

    # NFKC turns PDF ligatures (ﬁ, ﬂ) and full-width forms into plain letters; ₹ is unchanged.
    text = unicodedata.normalize("NFKC", text).replace("\r", "").replace("\u00a0", " ")
    text = re.sub(r"(\w)-\n(\w)", r"\1\2", text)
    text = re.sub(r"[ \t]+\n", "\n", text)
    text = re.sub(r"(?<!\n)\n(?!\n)", " ", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return re.sub(r"[ \t]{2,}", " ", text).strip()


def _split(text: str, limit: int = PASSAGE_CHARS) -> list[str]:
    """Split text into passages of at most ~limit chars on paragraph, then sentence, boundaries."""
    import re

    pieces: list[tuple[str, bool]] = []  # (text, starts a new paragraph)
    for para in (p.strip() for p in text.split("\n\n")):
        if not para:
            continue
        sentences = [para] if len(para) <= limit else re.split(r"(?<=[.!?])\s+", para)
        for k, sent in enumerate(sentences):
            while len(sent) > limit:  # a single very long "sentence" (tables, lists)
                pieces.append((sent[:limit], k == 0))
                sent, k = sent[limit:], 1
            pieces.append((sent, k == 0))
    passages: list[str] = []
    for piece, new_para in pieces:
        sep = "\n\n" if new_para else " "
        if passages and len(passages[-1]) + len(sep) + len(piece) <= limit:
            passages[-1] += sep + piece
        else:
            passages.append(piece)
    return passages


def _extract_pdf(data: bytes, name: str) -> tuple[list[dict], int]:
    import io

    from pypdf import PdfReader
    from pypdf.errors import PdfReadError

    try:
        reader = PdfReader(io.BytesIO(data))
        if reader.is_encrypted and not reader.decrypt(""):
            raise HTTPException(422, f"{name} is password-protected. Remove the password and upload it again.")
        pages = [_clean(page.extract_text() or "") for page in reader.pages]
    except PdfReadError as e:
        raise HTTPException(422, f"{name} doesn't look like a valid PDF ({e}).")
    passages = [
        {"text": chunk, "source": f"{name} · p.{i}"}
        for i, page in enumerate(pages, 1)
        for chunk in _split(page)
    ]
    if not passages:
        raise HTTPException(
            422,
            f"No selectable text found in {name}. It's probably a scanned image; run OCR on it first.",
        )
    return passages, len(pages)


def _table_to_sentences(table) -> str:
    """Write table rows as sentences. The detector reads "The price is $79." far more
    reliably than "Price | $79", which it tends to flag as unsupported."""
    rows = []
    for row in table.rows:
        cells = list(dict.fromkeys(c.text.strip() for c in row.cells))  # merged cells repeat
        cells = [c for c in cells if c]
        if cells:
            rows.append(cells)
    if not rows:
        return ""
    lower = lambda s: s[:1].lower() + s[1:] if s[:2] != s[:2].upper() else s  # keep acronyms
    if all(len(r) == 2 for r in rows):
        return " ".join(f"The {lower(k)} is {v.rstrip('.')}." for k, v in rows)
    header, body = rows[0], rows[1:] or rows
    lines = []
    for r in body:
        pairs = [f"{h} is {v.rstrip('.')}" for h, v in zip(header[1:], r[1:])]
        lines.append(f"{r[0]}: {'; '.join(pairs)}." if pairs else f"{' '.join(r)}.")
    return " ".join(lines)


def _extract_docx(data: bytes, name: str) -> tuple[list[dict], int]:
    import io
    import zipfile

    import docx

    try:
        document = docx.Document(io.BytesIO(data))
    except (zipfile.BadZipFile, KeyError, ValueError):
        raise HTTPException(422, f"{name} isn't a valid .docx file. Old .doc files must be saved as .docx first.")
    blocks = [p.text.strip() for p in document.paragraphs if p.text.strip()]
    for table in document.tables:
        blocks.append(_table_to_sentences(table))
    chunks = _split("\n\n".join(blocks))
    if not chunks:
        raise HTTPException(422, f"{name} has no text to extract.")
    return [{"text": c, "source": f"{name} · part {i}"} for i, c in enumerate(chunks, 1)], 0


@app.post("/api/extract")
async def extract(file: UploadFile = File(...)) -> dict:
    name = file.filename or "upload"
    data = await file.read(MAX_UPLOAD_BYTES + 1)
    if len(data) > MAX_UPLOAD_BYTES:
        raise HTTPException(413, f"{name} is larger than 20 MB.")
    if not data:
        raise HTTPException(422, f"{name} is empty.")
    ext = name.rsplit(".", 1)[-1].lower() if "." in name else ""
    # Decide by the file's actual contents, not only its name.
    is_pdf = data[:5] == b"%PDF-"
    is_zip = data[:4] == b"PK\x03\x04"
    is_ole = data[:8] == b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1"  # legacy .doc / .xls
    if is_pdf or (ext == "pdf" and not is_zip):
        passages, pages = _extract_pdf(data, name)
        kind = "pdf"
    elif is_zip and ext in ("docx", "docm", "dotx", ""):
        passages, pages = _extract_docx(data, name)
        kind = "docx"
    elif is_zip:
        raise HTTPException(415, f"{name} is a zip-based file ({ext or 'unknown type'}), not a Word document. Use PDF, DOCX, TXT or MD.")
    elif is_ole or ext == "doc":
        raise HTTPException(415, f"{name} is an old Word .doc file. Open it in Word and save it as .docx.")
    elif ext in ("txt", "md", "markdown", "text", ""):
        try:
            raw = data.decode("utf-8-sig")
        except UnicodeDecodeError:
            raw = data.decode("cp1252", errors="replace")  # common for Windows-saved .txt
        if reason := unreadable_reason(raw):
            raise HTTPException(422, f"{name} doesn't contain readable text ({reason}).")
        text = raw.strip() if ext in ("md", "markdown") else _clean(raw)
        passages = [{"text": c, "source": f"{name} · part {i}"} for i, c in enumerate(_split(text), 1)]
        pages, kind = 0, "text"
        if not passages:
            raise HTTPException(422, f"{name} has no text.")
    else:
        raise HTTPException(415, f"{name}: unsupported file type. Use PDF, DOCX, TXT or MD.")
    full_text = "\n\n".join(p["text"] for p in passages)
    if reason := unreadable_reason(full_text):  # e.g. a PDF whose fonts map to junk glyphs
        raise HTTPException(422, f"The text extracted from {name} is garbled ({reason}). Try exporting it again as PDF or DOCX.")
    return {"filename": name, "kind": kind, "pages": pages, "chars": len(full_text),
            "passages": passages, "text": full_text}


def find_free_port(host: str, start: int, attempts: int = 20) -> int:
    """Return ``start`` if it is free, otherwise the next free port after it."""
    import socket

    for port in range(start, start + attempts):
        with socket.socket(socket.AF_INET, socket.SOCK_STREAM) as sock:
            # Same option uvicorn uses, so a port lingering in TIME_WAIT after a restart
            # isn't mistaken for one that another server is listening on.
            sock.setsockopt(socket.SOL_SOCKET, socket.SO_REUSEADDR, 1)
            try:
                sock.bind((host, port))
                return port
            except OSError:
                continue
    raise SystemExit(f"No free port found between {start} and {start + attempts - 1}.")


if __name__ == "__main__":
    import uvicorn

    host = "127.0.0.1"
    wanted = int(os.environ.get("PORT", "8765"))
    port = find_free_port(host, wanted)
    if port != wanted:
        print(f"Port {wanted} is already in use (another GroundCheck may be running); using {port} instead.", flush=True)
    print(f"\n  GroundCheck is running at http://{host}:{port}\n", flush=True)
    uvicorn.run(app, host=host, port=port)
