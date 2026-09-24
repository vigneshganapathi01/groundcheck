"""Regression suite for GroundCheck's detection quality.

Runs every case against a running server and reports, per scoring mode, how many planted
errors were caught and how many correct facts were wrongly flagged.

    .venv/bin/python tests/regression_suite.py                      # local server
    GROUNDCHECK_URL=http://127.0.0.1:7860 .venv/bin/python tests/regression_suite.py

Exits non-zero if per-sentence mode (the default) raises any false alarm or catches fewer
errors than the recorded baseline.
"""

from __future__ import annotations

import io
import json
import os
import ssl
import sys
import urllib.request
import uuid

BASE = os.environ.get("GROUNDCHECK_URL", "http://127.0.0.1:8765").rstrip("/")
try:  # python.org builds on macOS ship without root certificates; use certifi's when present
    import certifi

    SSL = ssl.create_default_context(cafile=certifi.where())
except ImportError:
    SSL = None
BASELINE_CAUGHT = 11  # per-sentence mode, measured 2026-09-24 (torch and onnx backends)

TRAVEL_POLICY = [
    ("Employee Travel & Expense Policy", None),
    ("Travel Eligibility", "Employees traveling more than 50 kilometers from their assigned office for approved business purposes may claim eligible travel expenses."),
    ("Accommodation", "Hotel accommodation is reimbursable up to ₹6,000 per night. Employees should book hotels from the company's approved list wherever possible."),
    ("Meals", "Meal expenses are reimbursable up to ₹1,200 per day."),
    ("Transportation", "Economy-class airfare and standard train fares are reimbursable. Business-class airfare requires prior approval from the department head."),
    ("Expense Claims", "Expense claims must be submitted within 15 days of returning from the trip. Receipts are required for all expenses above ₹500."),
    ("Travel Advance", "Employees may request a travel advance up to 7 days before departure. Any unused advance must be returned within 10 days of returning from the trip."),
]
TRAVEL_ANSWER = (
    "Employees traveling more than 50 kilometers from their office for approved business purposes can claim travel "
    "expenses. Hotel accommodation is reimbursable up to ₹6,000 per night, and meals can be claimed up to ₹1,200 per "
    "day. Employees can claim economy-class airfare and standard train fares. Business-class airfare is also "
    "reimbursable for senior employees without additional approval. Expense claims must be submitted within 30 "
    "calendar days of returning from the trip, and receipts are required for expenses above ₹500. Employees may "
    "request a travel advance up to seven days before departure, and any unused advance must be returned within 10 "
    "days of returning."
)

# (name, context, question, answer, must be flagged, must NOT be flagged)
CASES = [
    ("Wrong number", ["France is a country in Europe. The capital of France is Paris. The population of France is 67 million."],
     "What is the capital of France? What is the population of France?",
     "The capital of France is Paris. The population of France is 69 million.", ["69 million"], ["capital of France is Paris"]),
    ("Invented fact", ["The Eiffel Tower was completed in 1889 for the World's Fair. It is 330 metres tall and is located on the Champ de Mars in Paris."],
     "Tell me about the Eiffel Tower.",
     "The Eiffel Tower was completed in 1889 and is 330 metres tall. It was designed by Leonardo da Vinci.", ["Leonardo da Vinci"], ["1889", "330 metres"]),
    ("Faithful (aspirin)", ["Aspirin is used to reduce fever and relieve mild to moderate pain. Adults typically take 300–600 mg every four to six hours, up to a maximum of 4 g per day."],
     "How much aspirin can an adult take per day?",
     "Adults can take up to a maximum of 4 g of aspirin per day, usually 300–600 mg every four to six hours.", [], ["4 g", "300–600 mg"]),
    ("Refund policy", ["Refund policy: Customers may return unused items within 30 days of delivery for a full refund. Items must be in their original packaging.",
                       "Shipping: Standard shipping is free on orders over $50. Return shipping costs are paid by the customer unless the item arrived damaged.",
                       "Electronics can be returned within 14 days and are subject to a 10% restocking fee."],
     "What is the return policy for a laptop I bought?",
     "You can return the laptop within 14 days of delivery, but a 10% restocking fee applies. Return shipping is always free, and you'll also receive a $20 store credit for the inconvenience.",
     ["$20 store credit", "always free"], ["14 days", "10% restocking fee"]),
    ("Context overrides facts", ["In this fictional world, the capital of France is Lyon."], "What is the capital of France?",
     "The capital of France is Paris.", ["Paris"], []),
    ("HR leave", ["Full-time employees receive 20 days of paid annual leave per calendar year. Unused leave of up to 5 days can be carried over to the next year.",
                  "Employees must submit leave requests at least two weeks in advance through the HR portal. Sick leave is separate and allows up to 10 days per year, with a doctor's note required for absences longer than 3 days."],
     "How many vacation days do I get, and can I carry them over?",
     "Full-time employees get 20 days of paid annual leave each year. You can carry over up to 10 unused days to the next year, and any remaining days are paid out in cash at the end of December.",
     ["paid out in cash", "10 unused days"], ["20 days of paid annual leave"]),
    ("Smartwatch faithful", ["The SmartBand X2 has a 1.4-inch AMOLED display and a battery that lasts up to 12 days on a single charge. It is water resistant up to 50 metres and supports heart-rate, sleep and blood-oxygen tracking. It charges fully in about 90 minutes using a magnetic cable."],
     "How long does the SmartBand X2 battery last, and is it waterproof?",
     "The SmartBand X2 battery lasts up to 12 days on a single charge, and the band is water resistant up to 50 metres.", [], ["12 days", "50 metres"]),
    ("Taj Mahal", ["The Taj Mahal is an ivory-white marble mausoleum in Agra, India. It was commissioned in 1631 by the Mughal emperor Shah Jahan to house the tomb of his wife Mumtaz Mahal. Construction of the main mausoleum was completed around 1643."],
     "Who built the Taj Mahal, and when?",
     "The Taj Mahal was commissioned in 1631 by the Mughal emperor Shah Jahan for his wife Mumtaz Mahal. More than 20,000 workers were employed on the project, and the main mausoleum was completed around 1648.",
     ["20,000 workers", "1648"], ["commissioned in 1631"]),
    ("Library summary", ["The city council approved a budget of $4.2 million for the new public library on Tuesday. Construction is expected to begin in March 2027 and finish by late 2028. The library will include a children's wing, a makerspace and 120 computer stations."],
     "", "The city council approved $4.2 million for a new public library. Construction starts in March 2027 and should finish by late 2028. The library will have a children's wing, a makerspace, 120 computer stations and a rooftop café.",
     ["rooftop café"], ["$4.2 million", "March 2027", "late 2028"]),
    ("Store hours", ["The store opens at 9 am and closes at 6 pm on weekdays. On Saturdays it opens from 10 am to 4 pm. The store is closed on Sundays and public holidays."],
     "Is the store open on Sunday?", "Yes, the store is open on Sundays from 10 am to 4 pm.", ["open on Sundays"], []),
]


def post_json(path: str, payload: dict) -> dict:
    req = urllib.request.Request(BASE + path, json.dumps(payload).encode(), {"Content-Type": "application/json"})
    return json.load(urllib.request.urlopen(req, timeout=900, context=SSL))


def upload(name: str, data: bytes) -> dict:
    boundary = uuid.uuid4().hex
    body = (f"--{boundary}\r\nContent-Disposition: form-data; name=\"file\"; filename=\"{name}\"\r\n"
            "Content-Type: application/octet-stream\r\n\r\n").encode() + data + f"\r\n--{boundary}--\r\n".encode()
    req = urllib.request.Request(BASE + "/api/extract", body, {"Content-Type": f"multipart/form-data; boundary={boundary}"})
    return json.load(urllib.request.urlopen(req, timeout=300, context=SSL))


def travel_policy_docx() -> bytes:
    import docx

    d = docx.Document()
    d.add_heading(TRAVEL_POLICY[0][0], 0)
    for heading, body in TRAVEL_POLICY[1:]:
        d.add_heading(heading, 1)
        d.add_paragraph(body)
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()


def covered(result: dict, answer: str, phrase: str) -> bool:
    i = answer.find(phrase)
    return any(s["start"] < i + len(phrase) and s["end"] > i for s in result["spans"])


def main() -> int:
    extracted = upload("travel_policy.docx", travel_policy_docx())
    assert "₹6,000" in extracted["text"] and "�" not in extracted["text"], "DOCX extraction is broken"
    print(f"DOCX extraction OK ({len(extracted['passages'])} passage(s), ₹ intact)")
    cases = [("Travel policy (DOCX)", [p["text"] for p in extracted["passages"]],
              "What are the rules for claiming travel expenses?", TRAVEL_ANSWER,
              ["reimbursable for senior employees", "30 calendar days"],
              ["₹6,000", "₹1,200", "economy-class", "₹500", "seven days", "10 days"])] + CASES

    totals = {}
    for mode in ("whole", "sentence"):
        caught = missed = false_alarms = 0
        print(f"\n=== mode: {mode}")
        for name, ctx, q, answer, must, must_not in cases:
            r = post_json("/api/detect", {"context": ctx, "question": q or None, "answer": answer,
                                          "threshold": 0.5, "mode": mode})
            hit = [m for m in must if covered(r, answer, m)]
            fp = [m for m in must_not if covered(r, answer, m)]
            caught, missed, false_alarms = caught + len(hit), missed + len(must) - len(hit), false_alarms + len(fp)
            note = (f"  missed {[m for m in must if m not in hit]}" if len(hit) < len(must) else "") + (f"  FALSE {fp}" if fp else "")
            print(f"  {name:26s} caught {len(hit)}/{len(must)}  false alarms {len(fp)}/{len(must_not)}{note}")
        totals[mode] = (caught, missed, false_alarms)
        print(f"  TOTAL caught {caught}, missed {missed}, false alarms {false_alarms}")

    caught, _, false_alarms = totals["sentence"]
    ok = false_alarms == 0 and caught >= BASELINE_CAUGHT
    print(f"\n{'PASS' if ok else 'FAIL'}: per-sentence mode caught {caught} (baseline {BASELINE_CAUGHT}), false alarms {false_alarms}")
    return 0 if ok else 1


if __name__ == "__main__":
    sys.exit(main())
